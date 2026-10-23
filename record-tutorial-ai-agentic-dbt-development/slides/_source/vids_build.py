"""Build the Google Vids package: slide PNGs, one PPTX per section (speaker notes = voiceover),
one Markdown script + recording cue sheet per section, and a README.

usage: python3 vids_build.py <slides_html_dir> <out_dir>
"""
import glob
import html
import os
import re
import sys

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(__file__))
import content_a  # noqa: E402
import content_b  # noqa: E402
import content_c  # noqa: E402
from record import REC  # noqa: E402
from tts import ALT, TTS  # noqa: E402

SRC, OUT = sys.argv[1], sys.argv[2]
IMG = os.path.join(OUT, '..', 'vids-images')
os.makedirs(IMG, exist_ok=True)
DECKS = content_a.DECKS + content_b.DECKS + content_c.DECKS
WPM = 150
EMU_PX = 12192000 / 1920


def plain(s):
    s = re.sub(r'<br\s*/?>', ' ', s or '')
    s = re.sub(r'<[^>]+>', '', s)
    return html.unescape(s).strip()


def md_inline(s):
    """slide HTML (from lib.md) back to Markdown-ish text for the .md files"""
    s = re.sub(r'<br\s*/?>', ' ', s or '')
    s = re.sub(r'<code>(.*?)</code>', lambda m: '`' + html.unescape(m.group(1)) + '`', s)
    s = re.sub(r'</?strong>', '**', s)
    s = re.sub(r'<[^>]+>', '', s)
    return html.unescape(s).strip()


# ------------------------------------------------------------------ render images
GRAB = r"""
() => {
  const s = document.querySelector('#canvas .slide.active');
  const out = [];
  s.querySelectorAll('.editable').forEach(e => {
    const v = e.closest('.val'), r = v.getBoundingClientRect();
    out.push({x: r.left, y: r.top, w: r.width, h: r.height, text: e.textContent.trim()});
    e.style.visibility = 'hidden';
  });
  return out;
}
"""
boxes = {}
with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={'width': 1920, 'height': 1080})
    for d in DECKS:
        code = 's%02d' % d['num']
        f = os.path.abspath(os.path.join(SRC, d['file'] + '.html'))
        pg.goto('file://' + f, wait_until='domcontentloaded')
        pg.add_style_tag(content='.slide{transition:none!important}#progress{display:none!important}')
        pg.wait_for_timeout(800)
        pg.keyboard.press('Home')
        for i in range(len(d['slides'])):
            if i:
                pg.keyboard.press('ArrowRight')
            pg.wait_for_timeout(90)
            bx = pg.evaluate(GRAB)
            if bx:
                boxes[(code, i)] = bx
            pg.screenshot(path=os.path.join(IMG, '%s-%02d.png' % (code, i + 1)))
            if bx:
                pg.evaluate("() => document.querySelectorAll('#canvas .slide.active .editable').forEach(e => e.style.visibility='')")
    b.close()

# ------------------------------------------------------------------ pptx + markdown
readme_rows = []
for d in DECKS:
    num, code = d['num'], 's%02d' % d['num']
    scripts = TTS[num]
    assert len(scripts) == len(d['slides']), (code, len(scripts), len(d['slides']))
    words = sum(len(s.split()) for s in scripts)
    tts_min = words / WPM
    rec_budget = max(0.0, 9.5 - tts_min)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i, sl in enumerate(d['slides']):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(IMG, '%s-%02d.png' % (code, i + 1)), 0, 0, prs.slide_width, prs.slide_height)
        for bx in boxes.get((code, i), []):
            tb = s.shapes.add_textbox(Emu(int(bx['x'] * EMU_PX)), Emu(int(bx['y'] * EMU_PX)),
                                      Emu(int(bx['w'] * EMU_PX)), Emu(int(bx['h'] * EMU_PX)))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.word_wrap = True
            tf.auto_size = None
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT
            r = tf.paragraphs[0].add_run()
            r.text = bx['text'] or '?'
            r.font.size = Pt(31)
            r.font.bold = True
            r.font.name = 'Inter'
            r.font.color.rgb = RGBColor(0xFB, 0xBF, 0x24)
        s.notes_slide.notes_text_frame.text = scripts[i]
    pptx_name = '%s-%s.pptx' % (code, d['file'][4:])
    prs.save(os.path.join(OUT, pptx_name))

    # prompts in slide order
    prompts = [plain(sl['text']) for sl in d['slides'] if sl['type'] == 'prompt']
    rec = REC[num]
    n = len(d['slides'])
    L = []
    L.append('# %s · %s: Google Vids script' % (code.upper(), plain(d['title'])))
    L.append('')
    L.append('**Import:** `%s` (speaker notes = the scripts below). **Voiceover:** %d words, about %.1f min at %d wpm. '
             '**Demo recording budget:** about %.1f min (Vids videos max 10 min).' % (pptx_name, words, tts_min, WPM, rec_budget))
    L.append('')
    L.append('**Scene order in Vids:** scenes 1 to %d (slides) → your demo recording → scene %d (recap). '
             'Drag the recap scene to the end after you add the recording.' % (n - 1, n))
    if any(k[0] == code for k in boxes):
        L.append('')
        L.append('**Before converting:** in Google Slides, replace the amber `?` numbers on the real-data slides with your dry-run results. '
                 'The scripts never read the numbers aloud, so only the slide changes.')
    L.append('')
    L.append('## Voiceover, scene by scene')
    L.append('')
    for i, sl in enumerate(d['slides']):
        t = plain(sl.get('title') or sl.get('text') or sl.get('page') or '')
        if sl['type'] == 'prompt':
            t = 'Prompt card'
        L.append('**Scene %d · %s**' % (i + 1, t[:80]))
        L.append('')
        L.append(scripts[i])
        L.append('')
        if (num, i + 1) in ALT:
            L.append('*If your dry run does not show this:* ' + ALT[(num, i + 1)])
            L.append('')
    L.append('## Demo recording: prompting and changing files')
    L.append('')
    L.append('**Set up before you hit record**')
    L.append('')
    for x in rec['setup']:
        L.append('- ' + x)
    L.append('')
    L.append('**Steps**')
    L.append('')
    for k, (act, pi, fil) in enumerate(rec['steps'], 1):
        line = '%d. %s' % (k, act)
        if fil:
            line += ' *File:* `%s`' % fil if '`' not in fil else ' *File:* %s' % fil
        L.append(line)
        if pi:
            L.append('')
            L.append('   > ' + prompts[pi - 1])
            L.append('')
    if rec.get('say'):
        L.append('')
        L.append('**Optional lines to say over the recording**')
        L.append('')
        for x in rec['say']:
            L.append('- "%s"' % x)
    if prompts:
        L.append('')
        L.append('**All prompts in this section, ready to paste**')
        L.append('')
        for k, pr in enumerate(prompts, 1):
            L.append('%d. %s' % (k, pr))
            L.append('')
    with open(os.path.join(OUT, '%s-%s.md' % (code, d['file'][4:])), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(L).rstrip() + '\n')
    readme_rows.append((code.upper(), plain(d['title']), n, tts_min, rec_budget, pptx_name))

R = ['# Google Vids package', '',
     'One video per section. Each section: slides with an AI voiceover first, then your screen recording of the demo in Rosetta dbt Studio, then the recap slide.', '',
     '| Section | Title | Scenes | Voiceover | Demo budget | Import file |', '|---|---|---|---|---|---|']
for c, t, n, tm, rb, pf in readme_rows:
    R.append('| %s | %s | %d | %.1f min | %.1f min | `%s` |' % (c, t, n, tm, rb, pf))
R += ['', 'Voiceover times assume about %d words per minute; check the first video and adjust.' % WPM, '',
      '## Workflow for one section', '',
      '1. **Numbers (real-data slides only).** Upload the `.pptx` to Google Drive and open it with Google Slides. Replace each amber `?` with your dry-run result.',
      '2. **Convert.** In Google Vids: new video → **Landscape** → **Convert Slides** → pick the presentation → select all slides. Keep **Include AI voiceover, script, background music and animation** on.',
      '3. **Script.** On the script step, click **Replace with speaker notes**. Each scene now has the voiceover from this package. Pick a voice, then import.',
      '4. **Record the demo.** Follow the cue sheet in the section\'s `.md` file. Record with Vids\' own recorder, or upload a screen recording. Trim the waiting while the agent works.',
      '5. **Order.** Drag the recap scene (the last slide) after the recording.',
      '6. **Check the length.** Vids videos are limited to 10 minutes. If the demo runs long, split it into a second video.', '',
      '## Notes', '',
      '- Scripts contain no square brackets on purpose: Vids reads `[ ]` as voice tags such as `[excitedly]`. Add tags yourself if you want them.',
      '- Code names (like `stg_tlc__yellow_trips`) are spoken as plain words ("the staging model for yellow trips"), because AI voices read underscores badly. The exact names are on the slides.',
      '- Regenerate this package after editing the slides: `python3 vids_build.py <slides folder> <out folder>` from the slides `_source` folder.']
with open(os.path.join(OUT, 'README.md'), 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(R) + '\n')
print('ok', len(DECKS), 'decks', len(glob.glob(os.path.join(IMG, '*.png'))), 'images', sum(len(v) for v in boxes.values()), 'number boxes')
